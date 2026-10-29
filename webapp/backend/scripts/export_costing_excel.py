"""
Export the live products catalog to an Excel workbook, one sheet per duration,
one column per product — same layout as the hand-maintained "Standard costing
sheet" reference.

Usage:
    python scripts/export_costing_excel.py

Writes exports/costing_export_{timestamp}.xlsx (repo root).
Also importable: build_costing_export() is used by the admin "Export Master
Costing Sheet" button in routers/costing.py.
"""
import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).parent.parent.parent.parent
sys.path.insert(0, str(ROOT / "webapp" / "backend"))

# (export_row_label, firebase_field_key) in display order.
# Margin/est-sales/per-kW rows are mapped by MEANING, not by the field's own
# (frozen, legacy) name — "Margin @10 % On Cost" actually holds 6% today,
# see the mass-update margin-drift investigation this export grew out of.
ROW_SPEC = [
    ("Battery Pack", "Battery Pack"),
    ("Voltage", "Voltage"),
    ("Ampheres capacity", "Ampheres capacity"),
    ("KW calculation", "KW calculation"),
    ("Cell Voltage", "Cell Voltage"),
    ("Cell Capacity", "Cell Capacity"),
    ("Combination of Cells in Series", "Combination of Cells in Series"),
    ("Combination of Cells in parallel", "Combination of Cells in parallel"),
    ("Total No Of Cells", "Total No Of Cells"),
    ("FOB Cost Of Cells", "FOB Cost Of Cells"),
    ("Total FOB Cost of Cells", "Total FOB Cost of Cells"),
    ("Clearing & Customs", "Clearing & Customs 1"),
    ("Total Landed cost In India", "Total Landed cost In India 1"),
    ("Cost In INR -(1)", "Cost In INR-(1)"),
    ("BMS/PCM", "BMS or PCM cost"),
    ("Clearing & Customs", "Clearing & Customs 2"),
    ("Total Landed cost In India", "Total Landed cost In India 2"),
    ("Cost In INR -(2)", "Cost In INR-(2)"),
    ("Cabinet (INR)", "Cabinet"),
    ("Bus Bar", "Bus Bar"),
    ("Holder/caps", "Holder or Caps"),
    ("Wire & Gasket & Other Accessories", "Wire & Gasket & Other Assesorries"),
    ("Terminals + Connectors", "Terminals+ Connectors"),
    ("MCB/Fuse", "MCB or Fuse"),
    ("Lugs & Slew", "Lugs & Slew"),
    ("Nut Bolts", "Nut Bolts"),
    ("Fiber glass + rod", "Fiber glass +rod"),
    ("Awg cables", "Awg cables"),
    ("Shipping Charges", "Shipping Charges"),
    ("Packaging cost with safety packs", "Packaging cost with safety packs"),
    ("Total Other Charges (3)", "Total Other Chargers(3)"),
    ("Landing cost of material (1+2+3)", "Landing cost of material (1+2+3)"),
    ("Production Labour & Assembly overheads", "Production Labour & Assembly overheads"),
    ("Warranty & Service provision", "Warranty  & Service provision"),
    ("Total Cost of Pack (A)", "Total Cost of Pack (A)"),
    ("Margin @6% On Cost", "Margin @10 % On Cost"),
    ("Estimated Sales Cost-(A+6)", "Estimated Sales Cost-(B)"),
    ("Margin @10% On Cost", "Margin @15% On Cost"),
    ("Estimated Sales Cost-(B)", "Estimated Sales Cost-(B+5)"),
    ("Per Kw Pricing @ cost (A)", "Per Kw Pricing @ cost (A)"),
    ("Per Kw pricing @ 1st level profit (A+6)", "Per Kw pricing @ ist level profit (B)"),
    ("Per Kw pricing @ 2nd level profit (B)", "Per Kw pricing @ 2nd evel profit (B+5)"),
    ("BMS/PCM", "BMS or PCM"),
    ("LFP/NCM", "Cell Chemistry"),
    ("Centre tap/non centre tap", "Centre tapping"),
    ("Cylindrical/Prismatic", "Type of Cell"),
    ("Application", "Application"),
    ("Soft pack/Metal enclosure", "Enclosure"),
    ("If Metal enclosure - Tower-type/rack-mountable", "Mount"),
    ("Brand and type of cell", "Brand and type of cell"),
    ("Installation indoor or outdoor", "Installation indoor or outdoor"),
    ("Battery Partcode", "Battery Partcode"),
    ("Dollar Rate", "Dollar Rate"),
    ("Creation Date", "Creation Date"),
    ("Created By", "Created By"),
]


def build_costing_export(active_only: bool = True) -> Path:
    import openpyxl
    from basefunctions import get_all_durations
    import pgfire

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for duration in get_all_durations():
        products = pgfire.get("products", duration)
        if not products:
            continue
        items = products.items() if isinstance(products, dict) else enumerate(products)
        rows = [p for _pid, p in items if isinstance(p, dict) and (not active_only or p.get("active"))]
        if not rows:
            continue
        rows.sort(key=lambda p: (str(p.get("Battery Pack", "")), str(p.get("Battery Partcode", ""))))

        ws = wb.create_sheet(title=duration[:31])
        for r_idx, (label, _key) in enumerate(ROW_SPEC, start=1):
            ws.cell(row=r_idx, column=1, value=label)
        for c_idx, product in enumerate(rows, start=2):
            for r_idx, (_label, key) in enumerate(ROW_SPEC, start=1):
                ws.cell(row=r_idx, column=c_idx, value=product.get(key, ""))

    out_dir = ROOT / "exports"
    out_dir.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    dest = out_dir / f"costing_export_{stamp}.xlsx"
    wb.save(dest)
    return dest


def main():
    dest = build_costing_export()
    print(f"Export written: {dest}")


if __name__ == "__main__":
    main()
